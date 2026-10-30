"""
Tests for storing formula results in generated workbooks.

Phone viewers show the stored result of a formula rather than recalculating,
so every formula in a download must carry its value.
"""

import io

import openpyxl
from django.http import HttpResponse

from core.xlsx_values import (
    XLSX_CONTENT_TYPE,
    XlsxCachedValuesMiddleware,
    fill_cached_values,
)


def _workbook_bytes():
    wb = openpyxl.Workbook()
    est = wb.active
    est.title = 'Estimate'
    datas = wb.create_sheet('Datas')
    datas['E8'], datas['I8'], datas['G8'] = 2, 36, 1
    datas['J8'] = '=SUM(E8*I8/G8)'
    est['B4'] = 12.5
    est['E4'] = '=Datas!J8'
    est['H4'] = '=ROUND(B4*E4,2)'
    est['H5'] = '=ROUND(H4*0.18,2)'
    est['D6'] = '=D99'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _values(data, sheet='Estimate'):
    return openpyxl.load_workbook(io.BytesIO(data), data_only=True)[sheet]


def test_formulas_get_stored_values():
    data = fill_cached_values(_workbook_bytes())
    est = _values(data)
    assert est['E4'].value == 72
    assert est['H4'].value == 900
    assert est['H5'].value == 162
    assert est['D6'].value == 0
    assert _values(data, 'Datas')['J8'].value == 72


def test_formulas_are_kept():
    data = fill_cached_values(_workbook_bytes())
    est = openpyxl.load_workbook(io.BytesIO(data))['Estimate']
    assert est['H4'].value == '=ROUND(B4*E4,2)'


def test_bad_input_is_returned_unchanged():
    assert fill_cached_values(b'not a workbook') == b'not a workbook'


def test_middleware_fills_xlsx_responses():
    original = _workbook_bytes()
    middleware = XlsxCachedValuesMiddleware(
        lambda request: HttpResponse(original, content_type=XLSX_CONTENT_TYPE)
    )
    response = middleware(None)
    assert _values(response.content)['H4'].value == 900
    assert response['Content-Length'] == str(len(response.content))


def test_middleware_ignores_other_responses():
    middleware = XlsxCachedValuesMiddleware(lambda request: HttpResponse(b'hello'))
    assert middleware(None).content == b'hello'
